#!/usr/bin/env python3
"""Generate the Wisdom Vault User FAQ document (.docx).

A plain-language companion to the Business Information document, written for
members and the general public who have questions about whether Wisdom Vault
is genuine, why there is a fee, and whether anyone is forced to pay.

Usage:  python3 scripts/generate_user_faq_docx.py
Output: Wisdom_Vault_User_FAQ.docx  (in the project root)
"""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

DARK = RGBColor(0x1F, 0x1F, 0x1F)
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GREY = RGBColor(0x59, 0x59, 0x59)
GOLD = RGBColor(0x9A, 0x6A, 0x00)

OUT = "Wisdom_Vault_User_FAQ.docx"

NAME = "The Vault"
EMAIL = "administration@wisdomvault.com"
WEBSITE = "https://wisdom-vault-peach.vercel.app"


def set_base_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = DARK
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, color, before, after in (
        ("Heading 1", 15, NAVY, 18, 8),
        ("Heading 2", 12.5, NAVY, 12, 6),
    ):
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)


def title_block(doc):
    t = doc.add_paragraph()
    r = t.add_run("Wisdom Vault")
    r.font.size = Pt(30)
    r.font.bold = True
    r.font.color.rgb = NAVY

    s = doc.add_paragraph()
    r = s.add_run("Questions and Answers for Our Members and the Public")
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = GOLD

    meta = doc.add_paragraph()
    r = meta.add_run(
        "A plain-language explanation of who we are, what you pay for, and why joining Wisdom "
        "Vault is always your own free choice.\nVersion 1.1  \u00b7  September 2026  \u00b7  Malawi"
    )
    r.font.size = Pt(10)
    r.font.color.rgb = GREY
    meta.paragraph_format.space_after = Pt(14)


def bullets(doc, items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        if isinstance(it, tuple):
            lead, rest = it
            r = p.add_run(lead)
            r.font.bold = True
            p.add_run(rest)
        else:
            p.add_run(it)


def body():
    doc = Document()
    for sec in doc.sections:
        sec.page_width = Inches(8.27)
        sec.page_height = Inches(11.69)
        sec.left_margin = Inches(0.9)
        sec.right_margin = Inches(0.9)
        sec.top_margin = Inches(0.9)
        sec.bottom_margin = Inches(0.9)

    set_base_styles(doc)
    title_block(doc)

    # Intro
    doc.add_paragraph(
        "We hear that some people are worried that Wisdom Vault might be a scam, or that members "
        "are being forced to pay. We understand why people ask, and we want to answer those "
        "questions directly and honestly. Please read this short document. If anything is still "
        "not clear, talk to us first \u2014 our contact details are at the end \u2014 and we will do "
        "our best to help."
    )

    # Q1
    doc.add_heading("1. Who runs Wisdom Vault?", level=1)
    doc.add_paragraph(
        f"Wisdom Vault is a digital learning platform operated in Malawi by {NAME}, an independent "
        "digital business. The business is not yet registered as a company, and we do not pretend "
        "otherwise. Our website is " + WEBSITE + " and our support email is listed at the end of "
        "this document. If you want to confirm who you are dealing with, contact us \u2014 we answer "
        "openly and we keep a verifiable record of every payment we receive."
    )

    # Q2
    doc.add_heading("2. What do people actually pay for?", level=1)
    doc.add_paragraph(
        "Wisdom Vault is a paid learning platform \u2014 like an online school or a magazine you "
        "subscribe to. When you pay the membership fee, you get access to learning content: "
        "courses, guides, ebooks and videos on topics such as business and finance, technology, "
        "personal growth and practical skills. The price is always shown to you before you pay."
    )
    doc.add_paragraph(
        "Your fee is payment for that content access. It is not an investment, it is not money "
        "we keep and promise to grow, and Wisdom Vault never promises that you will earn money "
        "back from it."
    )

    # Q3
    doc.add_heading("3. Is joining and paying by force?", level=1)
    doc.add_paragraph("No. Nobody is forced, pressured or tricked into joining or paying.")
    bullets(doc, [
        ("Registration is free and separate from payment. ", "You can create an account without "
         "paying anything."),
        ("You must first agree, on a clear consent screen, ", "that you understand this is a paid "
         "service and accept the Terms and Privacy Policy. The screen has an \u201cI Agree\u201d button and "
         "a \u201cDisagree\u201d button. If you click \u201cDisagree\u201d, the process stops and no account is made."),
        ("Every payment is authorized by you ", "on your own phone through your mobile-money "
         "provider (Airtel Money or TNM Mpamba via PayChangu). No one can deduct money from you "
         "without your own action."),
        ("There are no hidden automatic charges. ", "If there is ever a renewal fee, it is shown to "
         "you first and you pay it only if you choose to."),
    ])
    doc.add_paragraph(
        "If someone told you that you must join, or pressured you to pay before you understood "
        "what you were buying, please contact us immediately and we will review what happened. "
        "Remember that a payment can only be made by you, on your own phone, so never let anyone "
        "push you into sending money before you are sure. Please read section 6 before you pay."
    )

    # Q4
    doc.add_heading("4. What about the referral rewards \u2014 is this a pyramid scheme?", level=1)
    doc.add_paragraph(
        "No \u2014 and this is the easiest thing to check. In a pyramid scheme, people pay money mainly "
        "to get paid for bringing in more people, and most members end up losing. Wisdom Vault is "
        "different in important ways:"
    )
    bullets(doc, [
        ("You are buying content, not an opportunity. ", "Your fee pays for learning content that "
         "you can use immediately, whether or not you ever invite anyone."),
        ("Nobody must recruit anyone. ", "Members who never invite a single person keep full access "
         "to everything they paid for."),
        ("Rewards are bonuses, not income. ", "The referral reward is a fixed bonus paid only when "
         "a person you invited freely chooses to join and pay. It is never promised as a salary "
         "or a guaranteed income, and no one is promised they will get rich by inviting others."),
        ("The rules are the same for everyone. ", "There is one account per person, and rewards "
         "from fake or duplicate accounts are refused and reversed."),
    ])

    # Q5
    doc.add_heading("5. How do referral rewards work exactly?", level=1)
    doc.add_paragraph(
        "When you join, you receive your own referral link. If a friend uses it, joins willingly "
        "and pays the membership fee, and that payment is confirmed, you are rewarded: currently "
        "MWK 7,000 for a directly referred member, with a smaller secondary reward of MWK 1,000 "
        "where that level of the program applies. The amounts are published on the platform. "
        "Rewards are added to your wallet and can be withdrawn according to the withdrawal rules "
        "shown in the app."
    )

    # Q6
    doc.add_heading("6. What if I am not satisfied? Are fees refundable?", level=1)
    doc.add_paragraph(
        "Paid fees are non-refundable once you have paid. This is why the consent screen asks you "
        "to confirm, before your account is even created, that you understand Wisdom Vault is a "
        "paid membership service, and why the full amount is always shown before you pay. Please "
        "take your time, read this document and the Terms of Service, and only pay when you are "
        "sure you want to join."
    )
    bullets(doc, [
        ("Your earned balance always stays yours: ", "if a withdrawal cannot be paid out to you, the "
         "full amount is returned to your wallet automatically. Money that never left the platform "
         "always comes back to you."),
        ("You can see everything: ", "every payment and every reward appears in your wallet history "
         "with the payment reference."),
        ("You can leave at any time: ", "you can stop using the platform and ask us to close your "
         "account and delete your data. Contact us and we will help."),
    ])

    # Q7
    doc.add_heading("7. What should I do if I think something is wrong?", level=1)
    doc.add_paragraph(
        "First, contact us directly \u2014 most questions and complaints are resolved quickly when we "
        "can look at the account and the payment reference together. If you feel you were cheated "
        "or misled, tell us and we will investigate honestly. Wisdom Vault also cooperates with "
        "banks, mobile-money providers, the police and regulators, and it will never ask you to "
        "hide anything from them."
    )

    # Q8
    doc.add_heading("8. Why do some people join? A word about \u201cquick money\u201d", level=1)
    doc.add_paragraph(
        "People join Wisdom Vault for different reasons. Many join because they want the learning "
        "content. Others join mainly because they are in a hurry to make money \u2014 perhaps they saw "
        "the referral rewards, or someone showed them an earnings screenshot. If that is your "
        "situation, please read this section slowly before you pay anything."
    )
    bullets(doc, [
        ("There is no quick money here. ", "Wisdom Vault sells learning content. Referral rewards "
         "are small, depend entirely on other people freely choosing to join, and are never "
         "guaranteed. That is stated in our Terms of Service, not hidden in small print."),
        ("Anyone who promises you earnings is breaking our rules. ", "No member, agent or staff "
         "member may promise you a specific income, tell you that you will \u201cget rich\u201d, or push "
         "you to pay quickly. If anyone does that, report them to us with their name or link and "
         "we will act on it."),
        ("If you only want fast money, please do not join. ", "Join because you want the content "
         "and the skills it teaches. That is what your fee buys \u2014 it is the only thing we promise, "
         "and it is worth taking your time to decide."),
        ("Pressure is a red flag, always. ", "A genuine platform can wait for your decision. Anyone "
         "who says \u201cpay now or you lose the chance\u201d is not following our rules \u2014 and no legitimate "
         "opportunity disappears just because you took a day to think."),
    ])

    # Q9
    doc.add_heading("9. Is Wisdom Vault legal? What the law actually says", level=1)
    doc.add_paragraph(
        "Yes \u2014 selling access to digital learning content is ordinary, lawful commerce, the same "
        "as selling online courses, ebooks or a magazine subscription. But we know that simply "
        "saying \u201cit is legal\u201d is not enough, so here is what the law actually looks at when it "
        "asks whether a platform like this is legitimate."
    )
    doc.add_paragraph(
        "Malawian law prohibits pyramid and \u201cchain\u201d schemes \u2014 schemes where money mostly moves "
        "between participants and the only way to \u201cearn\u201d is to keep recruiting others. When "
        "regulators and courts test whether a business is such a scheme, the questions are "
        "practical ones like these:"
    )
    bullets(doc, [
        ("Is there a real product? ", "Yes. A member's fee buys real content \u2014 courses, guides, "
         "ebooks and videos \u2014 that they can use immediately. Recruiting is not the product; the "
         "content is."),
        ("Are people paid mainly for recruiting? ", "No. No member is required to recruit anyone, "
         "and members who never invite a single person keep everything they paid for. Referral "
         "rewards are a fixed bonus for real, valid sign-ups \u2014 never the main business."),
        ("Are returns promised? ", "No. We promise no earnings of any kind, and the Terms of "
         "Service say so expressly. Promising returns is exactly what makes a scheme illegal, "
         "which is why we forbid anyone from doing it on our behalf."),
        ("Is money taken honestly? ", "Every payment is made by you, on your own phone, after a "
         "consent screen and a clearly shown price. Nothing is ever deducted without your own "
         "action, and every transaction leaves a verifiable record with the payment gateway "
         "(PayChangu, Airtel Money or TNM Mpamba)."),
    ])
    doc.add_paragraph(
        "One thing we will never do is claim that a court has \u201capproved\u201d Wisdom Vault. No court "
        "has ruled on us, and any platform that waves around fake certificates, endorsements or "
        "\u201cgovernment approval\u201d should be treated with suspicion. What actually protects you is "
        "the law plus your own checks: read the Terms of Service, keep your payment references, "
        "and if you want an independent opinion, ask a lawyer, your bank or your mobile-money "
        "provider \u2014 we welcome those questions."
    )
    doc.add_paragraph(
        "That is also why we say it plainly: Wisdom Vault is not a fraud. We sell a real product, "
        "we never promise income, we never touch your money without your own authorization, we "
        "give you a full record of every transaction, and we cooperate with banks, mobile-money "
        "providers, PayChangu and law enforcement whenever a question is raised."
    )

    # Q10
    doc.add_heading("10. How can I reach you?", level=1)
    bullets(doc, [
        ("Support email: ", EMAIL),
        ("Website: ", WEBSITE),
        ("Operator: ", NAME + " (independent digital business, Malawi)"),
    ])
    doc.add_paragraph(
        "We encourage anyone with doubts to ask us before paying, and to report to us any person "
        "who claims they can force or guarantee anyone to earn money through Wisdom Vault \u2014 that "
        "is not how our platform works."
    )

    doc.add_paragraph()
    end = doc.add_paragraph()
    r = end.add_run(
        "This document is an explanation of Wisdom Vault's published terms and practices as of "
        "September 2026. The Terms of Service and Privacy Policy on our website are the documents "
        "that govern use of the platform."
    )
    r.font.size = Pt(9)
    r.font.italic = True
    r.font.color.rgb = GREY

    doc.save(OUT)
    print(f"Saved {OUT}")


if __name__ == "__main__":
    body()
